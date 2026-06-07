"""
Generates reference.docx for pandoc with BTEC formatting:
  - Times New Roman 12pt body
  - 1.5 line spacing
  - A4 page
  - Margins: Left 3cm, Right 1.5cm, Top 2cm, Bottom 2cm
  - Justified text, first-line indent 1.25cm
  - Page numbers in bottom center
"""
from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement


def set_default_font(doc):
    styles = doc.styles
    style = styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    # Make sure East Asian and complex script slots also use Times New Roman
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rFonts.set(qn(f'w:{attr}'), 'Times New Roman')

    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)


def set_heading_styles(doc):
    sizes = {'Heading 1': 16, 'Heading 2': 14, 'Heading 3': 12, 'Heading 4': 12}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        pf = style.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.5
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)


def set_page(doc):
    section = doc.sections[0]
    section.page_height = Mm(297)   # A4
    section.page_width = Mm(210)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number_field(p)


def main():
    doc = Document()
    set_page(doc)
    set_default_font(doc)
    set_heading_styles(doc)
    add_footer(doc)
    # A placeholder paragraph so pandoc copies the style
    doc.add_paragraph("placeholder")
    doc.save('reference.docx')
    print("reference.docx written")


if __name__ == '__main__':
    main()
